"""Phase 2 mode tests: learn -> make -> outline -> blend -> restyle -> edit -> critique."""

from __future__ import annotations

import json

import pytest

from deckforge_core.modes import (
    ModeError,
    parse_edit,
    run_blend,
    run_critique,
    run_edit,
    run_learn,
    run_make,
    run_outline,
    run_restyle,
)
from deckforge_core.modes.common import list_packs, resolve_pack
from deckforge_core.modes.packs import cmd_delete
from deckforge_core.tools.generate_corpus import build_deck


@pytest.fixture()
def corpus(tmp_path):
    out = tmp_path / "corpus"
    out.mkdir()
    for i in range(3):
        build_deck(out / f"deck_{i:02d}.pptx", i, seed=20240316)
    return out


@pytest.fixture()
def packed(tmp_path, corpus, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    result = run_learn(corpus, "demo")
    assert result.decks == 3
    return result.pack_dir


def test_learn_and_resolve(tmp_path, packed, monkeypatch):
    assert list_packs() == ["demo"]
    pack, library = resolve_pack("demo")
    assert pack.name == "demo"
    assert len(pack.archetypes) >= 5
    assert library is not None
    assert (packed / "pack.json").is_file()
    assert (packed / "style_profile.json").is_file()


def test_learn_missing_source_raises(tmp_path, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    with pytest.raises(ModeError):
        run_learn(tmp_path / "nope", "demo")


def test_learn_empty_folder_raises(tmp_path, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(ModeError):
        run_learn(empty, "demo")


def test_make_renders_pptx(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    out = tmp_path / "made"
    result = run_make("Product strategy", "demo", str(out), slides=8)
    assert result.slide_count == 8
    assert result.out_path.endswith(".pptx")
    assert (tmp_path / "made.pptx").is_file()


def test_make_saves_plan(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    plan_json = tmp_path / "plan.json"
    run_make("Metrics update", "demo", str(tmp_path / "m2"), slides=5, save_plan=str(plan_json))
    assert plan_json.is_file()
    data = json.loads(plan_json.read_text(encoding="utf-8"))
    assert len(data["slides"]) == 5


def test_make_requires_prompt(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    with pytest.raises(ModeError):
        run_make("   ", "demo", str(tmp_path / "x"))


def test_make_missing_pack_raises(tmp_path, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    with pytest.raises(ModeError):
        run_make("something", "not-a-pack", str(tmp_path / "x"))


@pytest.mark.parametrize(
    ("aspect", "expected_landscape"),
    [("16:9", True), ("4:3", True), ("portrait", False), ("9:16", False)],
)
def test_make_aspect_ratios(tmp_path, packed, monkeypatch, aspect, expected_landscape):
    from pptx import Presentation
    from pptx.util import Emu

    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    out = tmp_path / (aspect.replace(":", "-"))
    result = run_make(
        "Aspect check", "demo", str(out), slides=6, aspect=aspect
    )
    assert result.out_path.endswith(".pptx")
    presentation = Presentation(result.out_path)
    landscape = Emu(presentation.slide_width) > Emu(presentation.slide_height)
    assert landscape is expected_landscape


def test_outline_matches_prompt_length(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    lines = run_outline("Quarterly review", "demo", slides=6)
    assert len(lines) == 6
    assert all(1 <= line.n <= 6 for line in lines)
    assert all(line.archetype for line in lines)


def test_restyle_render(tmp_path, packed, corpus, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    out = tmp_path / "restyled"
    source = next(corpus.glob("*.pptx"))
    result = run_restyle(source, "demo", out)
    assert result.slide_count >= 1
    assert (tmp_path / "restyled.pptx").is_file()


def test_critique_offline(tmp_path, packed, corpus, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    source = next(corpus.glob("*.pptx"))
    report = run_critique(source, "demo", out_dir=tmp_path / "qa")
    assert report.rendered is False
    assert isinstance(report.slides, list)


def test_blend_two_packs(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    result = run_blend("demo", "demo", "fused")
    assert result.name == "fused"
    assert result.archetypes
    pack, _ = resolve_pack("fused")
    assert pack.source_deck_count == 6
    cmd_delete("fused")
    assert list_packs() == ["demo"]


def test_blend_forbids_existing_name(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    occupied = tmp_path / "packs" / "occupied"
    occupied.mkdir(parents=True)
    (occupied / "pack.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ModeError):
        run_blend("demo", "demo", "occupied")


def test_edit_apply_and_render(tmp_path, packed, monkeypatch):
    monkeypatch.setenv("DECKFORGE_DATA_DIR", str(tmp_path))
    plan_json = tmp_path / "plan.json"
    run_make("Launch brief", "demo", str(tmp_path / "m3"), slides=4, save_plan=str(plan_json))
    out = tmp_path / "edited"
    edits = ["1:subtitle=A sharp subtitle", "2:items=One|Two"]
    result = run_edit(plan_json, edits, out, pack_spec="demo")
    assert result.applied == 2
    assert result.render.slide_count == 4
    assert (tmp_path / "edited.pptx").is_file()


def test_parse_edit_grammar():
    spec = parse_edit("3:card_1.paragraphs=Alpha|Beta")
    assert spec.slide == 3
    assert spec.slot == "card_1"
    assert spec.field == "paragraphs"
    assert spec.value == "Alpha|Beta"

    spec = parse_edit("2:title=Hello world")
    assert spec.field == "text"
    assert spec.value == "Hello world"


def test_parse_edit_rejects_bad_grammar():
    with pytest.raises(ModeError):
        parse_edit("no-equals")
    with pytest.raises(ModeError):
        parse_edit("x:title=hi")
    with pytest.raises(ModeError):
        parse_edit("2=hi")
