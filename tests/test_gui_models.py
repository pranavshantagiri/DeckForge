"""Model layer behind the GUI: pack learning, generation, previews, desktop glue.

These tests never import PySide6, so they also cover the GUI in CI jobs that do
not install the ``gui`` extra.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from deckforge_core.config import Settings
from deckforge_core.gui import models
from deckforge_core.providers import NoneRenderer
from deckforge_core.providers.render import RenderBackend, RenderedSlide
from deckforge_core.renderer.pack_io import export_dfpack, load_pack
from deckforge_core.schemas.deck_plan import DeckPlan

REPO_ROOT = Path(__file__).resolve().parents[1]
CORPUS = REPO_ROOT / "corpus" / "raw"


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="module")
def source_folder(tmp_path_factory) -> Path:
    decks = sorted(CORPUS.glob("*.pptx"))[:2]
    if not decks:
        pytest.skip("corpus/raw has no decks; cannot learn a pack")
    folder = tmp_path_factory.mktemp("source-decks")
    for deck in decks:
        shutil.copy2(deck, folder / deck.name)
    return folder


@pytest.fixture(scope="module")
def learned(source_folder, tmp_path_factory) -> models.LearnPackResult:
    packs = tmp_path_factory.mktemp("packs")
    return models.learn_pack_from_folder(
        source_folder,
        "gui-pack",
        target_dir=packs / "gui-pack",
        use_cache=False,
    )


@pytest.fixture
def out_dir(tmp_path) -> Path:
    return tmp_path / "out"


def _request(pack_dir: Path, out_dir: Path, **kwargs) -> models.GenerationRequest:
    payload = {
        "pack_dir": pack_dir,
        "brief": "Q3 results for the leadership team",
        "slide_count": 6,
        "render_backend": "none",
        "out_dir": out_dir,
    }
    payload.update(kwargs)
    return models.GenerationRequest(**payload)


# --------------------------------------------------------------------------- #
# Fake slide renderer (keeps COM out of the tests)
# --------------------------------------------------------------------------- #
class _FakeRenderer:
    backend = RenderBackend.POWERPOINT_COM
    backend_value = RenderBackend.POWERPOINT_COM.value

    def __init__(self, missing_from: int = 0) -> None:
        self.missing_from = missing_from

    def render(self, pptx_path, out_dir, *, slides=None, dpi=96):
        target = Path(out_dir)
        target.mkdir(parents=True, exist_ok=True)
        count = len(Presentation(str(pptx_path)).slides)
        results = []
        for number in range(1, count + 1):
            png = target / f"slide_{number:03d}.png"
            if number >= self.missing_from > 0:
                results.append(RenderedSlide(number=number, image_path="", error="no export"))
                continue
            Image.new("RGB", (320, 180), (20 * number % 255, 90, 140)).save(png, "PNG")
            results.append(
                RenderedSlide(number=number, image_path=str(png), width_px=320, height_px=180)
            )
        return results


class _FakeAutodetect:
    """Stands in for RendererAutodetect so no COM/LibreOffice call happens."""

    renderer_class = _FakeRenderer

    def __init__(self) -> None:
        self.com = True
        self.libreoffice = False

    def resolve(self, preferred: str = "auto"):
        if preferred == RenderBackend.NONE.value:
            return NoneRenderer()
        return type(self).renderer_class()


class _PartialAutodetect(_FakeAutodetect):
    renderer_class = None

    def resolve(self, preferred: str = "auto"):
        if preferred == RenderBackend.NONE.value:
            return NoneRenderer()
        return _FakeRenderer(missing_from=3)


@pytest.fixture
def fake_renderer(monkeypatch):
    monkeypatch.setattr(models, "RendererAutodetect", _FakeAutodetect)
    return _FakeAutodetect


# --------------------------------------------------------------------------- #
# Locations and helpers
# --------------------------------------------------------------------------- #
def test_slugify_and_roots(tmp_path):
    assert models.slugify("Q3 Results: Growth & Margin!") == "q3-results-growth-margin"
    assert models.slugify("   ", fallback="pack") == "pack"
    packs = tmp_path / "packs"
    assert models.packs_root(packs) == packs
    decks = models.decks_root(tmp_path / "decks")
    assert decks.is_dir()


def test_decks_in_finds_nested_pptx(tmp_path):
    (tmp_path / "a.pptx").write_bytes(b"x")
    (tmp_path / "nested").mkdir()
    (tmp_path / "nested" / "b.pptx").write_bytes(b"x")
    (tmp_path / "notes.txt").write_text("ignore me", encoding="utf-8")
    found = models.decks_in(tmp_path)
    assert [path.name for path in found] == ["a.pptx", "b.pptx"]
    assert models.decks_in(tmp_path / "missing") == []


def test_settings_file_points_at_the_json(tmp_path):
    settings = Settings(root=tmp_path)
    assert models.settings_file(settings) == tmp_path / Settings.FILE_NAME


def test_provider_and_backend_registries():
    assert "mock" in models.available_llm_providers()
    assert "local" in models.secret_providers()
    backends = models.available_render_backends()
    known = {backend.value for backend in RenderBackend}
    assert backends
    assert set(backends) <= known


# --------------------------------------------------------------------------- #
# Learn
# --------------------------------------------------------------------------- #
def test_learn_pack_writes_pack_and_reports_progress(source_folder, tmp_path):
    steps: list[tuple[int, str]] = []
    result = models.learn_pack_from_folder(
        source_folder,
        "Stepped Pack",
        target_dir=tmp_path / "packs" / "stepped-pack",
        progress=lambda percent, message: steps.append((percent, message)),
        use_cache=False,
    )
    assert (result.pack_dir / "pack.json").is_file()
    assert (result.pack_dir / "blueprints").is_dir()
    assert result.pack.name == "Stepped Pack"
    assert result.deck_count == 2
    assert result.slide_count > 5
    assert result.archetypes
    assert result.ingest_errors == []
    assert result.summary
    assert steps[0][0] < steps[-1][0] <= 95
    assert any("Reading" in message for _percent, message in steps)


def test_learn_pack_rejects_folder_without_decks(tmp_path):
    (tmp_path / "readme.txt").write_text("no decks here", encoding="utf-8")
    with pytest.raises(ValueError, match="no .pptx decks"):
        models.learn_pack_from_folder(
            tmp_path, "empty", target_dir=tmp_path / "out", use_cache=False
        )


def test_learn_pack_reports_unreadable_decks(tmp_path):
    (tmp_path / "broken.pptx").write_bytes(b"this is not a powerpoint file")
    with pytest.raises(ValueError, match="no readable decks"):
        models.learn_pack_from_folder(
            tmp_path, "broken", target_dir=tmp_path / "out", use_cache=False
        )


def test_learn_pack_records_ingest_errors_but_keeps_going(source_folder, tmp_path):
    (source_folder / "broken.pptx").write_bytes(b"not a deck")
    try:
        result = models.learn_pack_from_folder(
            source_folder, "mixed", target_dir=tmp_path / "out", use_cache=False
        )
    finally:
        (source_folder / "broken.pptx").unlink()
    assert result.deck_count == 2
    assert len(result.ingest_errors) == 1
    assert "broken.pptx" in result.ingest_errors[0]


# --------------------------------------------------------------------------- #
# Packs home
# --------------------------------------------------------------------------- #
def test_list_pack_entries_reads_saved_packs(learned):
    entries = models.list_pack_entries(learned.pack_dir.parent)
    assert [entry.name for entry in entries] == ["gui-pack"]
    entry = entries[0]
    assert entry.usable
    assert entry.deck_count == 2
    assert entry.slide_count == learned.slide_count
    assert entry.aspect_ratios
    assert "decks" in entry.summary
    assert entry.path == learned.pack_dir


def test_list_pack_entries_flags_unreadable_pack(tmp_path):
    (tmp_path / "broken").mkdir()
    (tmp_path / "broken" / "pack.json").write_text("{}", encoding="utf-8")
    entries = models.list_pack_entries(tmp_path)
    assert len(entries) == 1
    assert not entries[0].usable
    assert "unreadable" in entries[0].summary


def test_load_pack_from_dir_round_trips(learned):
    pack, library = models.load_pack_from_dir(learned.pack_dir)
    assert pack.name == "gui-pack"
    assert library.archetypes()


def test_import_pack_zip_extracts_dfpack(learned, tmp_path):
    archive = tmp_path / "gui-pack.dfpack"
    export_dfpack(learned.pack_dir, archive)
    destination = tmp_path / "imported"
    entry = models.import_pack_zip(archive, destination)
    assert entry.name == "gui-pack"
    assert (entry.path / "pack.json").is_file()
    assert [item.name for item in models.list_pack_entries(destination)] == ["gui-pack"]


def test_import_pack_zip_avoids_collisions(learned, tmp_path):
    archive = tmp_path / "gui-pack.dfpack"
    export_dfpack(learned.pack_dir, archive)
    destination = tmp_path / "imported"
    first = models.import_pack_zip(archive, destination)
    second = models.import_pack_zip(archive, destination)
    assert first.path != second.path
    assert len(list(models.list_pack_entries(destination))) == 2


# --------------------------------------------------------------------------- #
# Generation
# --------------------------------------------------------------------------- #
def test_generate_offline_renders_parseable_pptx(learned, out_dir):
    result = models.generate(_request(learned.pack_dir, out_dir))
    assert result.pptx_path.is_file()
    assert result.pptx_path.parent == out_dir
    assert result.pptx_path.suffix == ".pptx"
    presentation = Presentation(str(result.pptx_path))
    assert len(presentation.slides) == result.slide_count == 6
    assert all(slide.shapes for slide in presentation.slides)
    assert result.plan.pack == "gui-pack"
    assert result.qa is not None
    assert result.errors_total >= 0
    assert len(result.previews) == result.slide_count
    assert result.rendered_previews == 0
    assert result.backend == RenderBackend.NONE.value


def test_generate_target_pptx_is_slugged(learned, out_dir):
    request = _request(learned.pack_dir, out_dir, brief="Renewables: 2026 Outlook!")
    assert request.target_pptx().name == "renewables-2026-outlook.pptx"


def test_generate_progress_reaches_the_end(learned, out_dir):
    steps: list[tuple[int, str]] = []
    result = models.generate(
        _request(learned.pack_dir, out_dir), progress=lambda p, m: steps.append((p, m))
    )
    assert steps[-1][0] == 100
    assert str(result.pptx_path.name) in steps[-1][1]
    assert [percent for percent, _ in steps] == sorted(percent for percent, _ in steps)


def test_generate_requires_a_brief(learned, out_dir):
    with pytest.raises(ValueError, match="brief"):
        models.generate(_request(learned.pack_dir, out_dir, brief="   "))


def test_generate_fails_on_unknown_pack_dir(out_dir):
    request = _request(out_dir / "does-not-exist", out_dir)
    with pytest.raises(Exception):
        models.generate(request)


def test_generate_can_skip_qa_checks(learned, out_dir):
    result = models.generate(
        _request(learned.pack_dir, out_dir), run_checks=False
    )
    assert result.qa is None
    assert result.errors_total == 0
    assert all(not preview.issues for preview in result.previews)


def test_generate_uses_new_deck_dir_when_out_dir_missing(learned, tmp_path, monkeypatch):
    monkeypatch.setattr(models, "decks_root", lambda root=None: tmp_path / "decks")
    request = _request(learned.pack_dir, Path())
    request.out_dir = None
    result = models.generate(request)
    assert result.pptx_path.parent.parent == tmp_path / "decks"


def test_build_plan_offline_is_deterministic(learned):
    pack, _library = models.load_pack_from_dir(learned.pack_dir)
    request = _request(learned.pack_dir, learned.pack_dir)
    first = models.build_plan(request, pack)
    second = models.build_plan(request, pack)
    assert isinstance(first, DeckPlan)
    assert first.created_by == "template-planner"
    assert first.model_dump(mode="json") == second.model_dump(mode="json")
    assert first.slides[0].archetype == "title"
    assert first.slides[-1].archetype == "closing"


def test_build_plan_with_llm_falls_back_cleanly(learned, tmp_path):
    pack, _library = models.load_pack_from_dir(learned.pack_dir)
    request = _request(learned.pack_dir, learned.pack_dir, use_llm=True, llm_provider="mock")
    plan = models.build_plan(request, pack, settings=Settings(root=tmp_path))
    assert plan.slides
    assert plan.created_by == "template-planner"
    assert models.last_warnings


# --------------------------------------------------------------------------- #
# Previews
# --------------------------------------------------------------------------- #
def test_render_slide_previews_returns_existing_images(learned, out_dir, fake_renderer):
    result = models.generate(_request(learned.pack_dir, out_dir, render_backend="auto"))
    assert result.backend == RenderBackend.POWERPOINT_COM.value
    assert result.rendered_previews == result.slide_count
    for preview in result.previews:
        assert Path(preview.image_path).is_file()
        assert not preview.error


def test_render_slide_previews_skips_failed_slides(learned, out_dir, monkeypatch):
    monkeypatch.setattr(models, "RendererAutodetect", _PartialAutodetect)
    result = models.generate(_request(learned.pack_dir, out_dir, render_backend="auto"))
    assert result.rendered_previews == 2
    assert len(result.previews) == result.slide_count
    assert any("slide 3" in warning for warning in result.warnings)


def test_render_slide_previews_disabled_returns_nothing(learned, out_dir):
    backend, images, notes = models.render_slide_previews(
        out_dir / "missing.pptx", out_dir / "renders", "none"
    )
    assert backend == RenderBackend.NONE.value
    assert images == {}
    assert notes


def test_render_slide_previews_reports_unavailable_backend(learned, out_dir, monkeypatch):
    class _Missing:
        def resolve(self, preferred="auto"):
            raise RuntimeError("not installed")

    monkeypatch.setattr(models, "RendererAutodetect", lambda: _Missing())
    backend, images, notes = models.render_slide_previews(
        out_dir / "missing.pptx", out_dir / "renders", "powerpoint-com"
    )
    assert backend == RenderBackend.NONE.value
    assert images == {}
    assert "not installed" in notes[0]


# --------------------------------------------------------------------------- #
# Desktop integration
# --------------------------------------------------------------------------- #
def test_open_in_default_app_uses_startfile(tmp_path, monkeypatch):
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"x")
    calls: list[str] = []
    monkeypatch.setattr(os, "startfile", calls.append, raising=False)
    assert models.open_in_default_app(deck) == str(deck)
    assert calls == [str(deck)]


def test_open_in_default_app_rejects_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        models.open_in_default_app(tmp_path / "nope.pptx")


def test_reveal_in_file_manager_selects_the_file(tmp_path, monkeypatch):
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"x")
    calls: list[list[str]] = []
    monkeypatch.setattr(subprocess, "Popen", lambda args, *a, **k: calls.append(args))
    assert models.reveal_in_file_manager(deck) == str(deck.resolve())
    assert calls
    assert calls[0][0] in ("explorer", "open", "xdg-open")
    with pytest.raises(FileNotFoundError):
        models.reveal_in_file_manager(tmp_path / "gone.pptx")


# --------------------------------------------------------------------------- #
# Persistence round-trip through the pack writer the GUI relies on
# --------------------------------------------------------------------------- #
def test_learned_pack_is_valid_json_on_disk(learned):
    payload = json.loads((learned.pack_dir / "pack.json").read_text(encoding="utf-8"))
    assert payload["name"] == "gui-pack"
    assert payload["format_version"] == 1
    pack, _library = load_pack(learned.pack_dir)
    assert pack.name == payload["name"]
