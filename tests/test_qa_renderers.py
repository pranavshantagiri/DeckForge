"""QA renderer wrappers (Workstream H): COM + LibreOffice backends."""

from __future__ import annotations

import os
from pathlib import Path

import pytest

import deckforge_core.qa  # noqa: F401  (registers the concrete renderers)
from deckforge_core.providers import RENDERER_CLASSES
from deckforge_core.providers.render import (
    RenderBackend,
    RendererAutodetect,
    libreoffice_available,
    powerpoint_com_available,
)
from deckforge_core.qa.renderers import LibreOfficeRenderer, PowerPointCOMRenderer


def _make_pptx(path: Path, n_slides: int = 2) -> Path:
    from pptx import Presentation

    prs = Presentation()
    for index in range(n_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(500000, 500000, 3000000, 1000000)
        box.text_frame.text = f"QA renderer slide {index + 1}"
    prs.save(str(path))
    return path


def test_powerpoint_backend_identity():
    assert PowerPointCOMRenderer().backend == RenderBackend.POWERPOINT_COM
    assert PowerPointCOMRenderer.name == "powerpoint-com"


def test_libreoffice_backend_identity():
    assert LibreOfficeRenderer().backend == RenderBackend.LIBREOFFICE
    assert LibreOfficeRenderer.name == "libreoffice"


def test_qa_import_registers_renderers():
    assert RENDERER_CLASSES[RenderBackend.POWERPOINT_COM] is PowerPointCOMRenderer
    assert RENDERER_CLASSES[RenderBackend.LIBREOFFICE] is LibreOfficeRenderer
    assert RenderBackend.NONE in RENDERER_CLASSES  # NoneRenderer still registered


def test_detect_reports_current_machine():
    assert PowerPointCOMRenderer.available() is powerpoint_com_available()
    assert LibreOfficeRenderer.available() is libreoffice_available()


def test_resolve_auto_prefers_installed_backend():
    detector = RendererAutodetect()
    resolved = detector.resolve(RenderBackend.AUTO.value)
    if detector.com:
        assert isinstance(resolved, PowerPointCOMRenderer)
    elif detector.libreoffice:
        assert isinstance(resolved, LibreOfficeRenderer)
    else:
        from deckforge_core.providers._registry import NoneRenderer

        assert isinstance(resolved, NoneRenderer)


def test_resolve_powerpoint_com_when_available():
    if not RendererAutodetect().com:
        pytest.skip("PowerPoint COM is not available on this machine")
    resolved = RendererAutodetect().resolve(RenderBackend.POWERPOINT_COM.value)
    assert isinstance(resolved, PowerPointCOMRenderer)


def test_resolve_libreoffice_when_available():
    if not RendererAutodetect().libreoffice:
        pytest.skip("LibreOffice is not available on this machine")
    resolved = RendererAutodetect().resolve(RenderBackend.LIBREOFFICE.value)
    assert isinstance(resolved, LibreOfficeRenderer)


@pytest.mark.skipif(
    not powerpoint_com_available(), reason="PowerPoint COM not available on this machine"
)
@pytest.mark.skipif(
    os.environ.get("DECKFORGE_TEST_COM_RENDER") != "1",
    reason="set DECKFORGE_TEST_COM_RENDER=1 to exercise live PowerPoint COM",
)
def test_powerpoint_com_renders_pngs(tmp_path):
    pptx = _make_pptx(tmp_path / "tiny.pptx", n_slides=2)
    images = PowerPointCOMRenderer().render(pptx, tmp_path / "renders")
    assert [r.number for r in images] == [1, 2]
    for result in images:
        assert result.error is None
        assert result.width_px > 0 and result.height_px > 0
        assert Path(result.image_path).exists()
        assert Path(result.image_path).stat().st_size > 0


@pytest.mark.skipif(
    powerpoint_com_available(), reason="PowerPoint COM IS available; graceful path not exercised"
)
def test_powerpoint_com_render_degrades_without_pywin32(tmp_path):
    pptx = _make_pptx(tmp_path / "tiny.pptx", n_slides=3)
    results = PowerPointCOMRenderer().render(pptx, tmp_path / "renders")
    assert [r.number for r in results] == [1, 2, 3]
    assert all(r.error for r in results)


@pytest.mark.skipif(
    libreoffice_available(), reason="LibreOffice IS available; graceful path not exercised"
)
def test_libreoffice_render_degrades_when_absent(tmp_path):
    pptx = _make_pptx(tmp_path / "tiny.pptx", n_slides=2)
    results = LibreOfficeRenderer().render(pptx, tmp_path / "renders")
    assert [r.number for r in results] == [1, 2]
    assert all(r.error for r in results)
