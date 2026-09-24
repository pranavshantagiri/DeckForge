"""Render tests: real python-pptx output, reopen round-trip, editable text,
connector styles, and graceful no-throw behaviour on bad specs.

Elbow routes have 4 points (down/across/down) so each rendered edge emits 3
chained straight connectors; the round-trip test pins that count.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.shapes.connector import Connector
from pptx.util import Emu, Pt

from deckforge_core.diagrams import render_diagram
from deckforge_core.schemas.blueprints import SlotRegion
from deckforge_core.schemas.deck_plan import DiagramSpec, GraphEdge, GraphNode, GraphSpec
from deckforge_core.schemas.pack import (
    ColorPalette,
    FontPair,
    Grid,
    Margins,
    PaletteColor,
    StyleProfile,
    TypeScale,
    TypeScaleEntry,
)

SLIDE_W, SLIDE_H = 12192000, 6858000
CANVAS = (SLIDE_W, SLIDE_H)
REGION = SlotRegion(x=0.08, y=0.12, w=0.84, h=0.7)


def _style() -> StyleProfile:
    return StyleProfile(
        palette=ColorPalette(
            colors=[
                PaletteColor(role="bg", hex="#FAFAF7", usage_pct=30.0),
                PaletteColor(role="text", hex="#1C1C1A", usage_pct=20.0),
                PaletteColor(role="muted-text", hex="#6B6B66", usage_pct=10.0),
                PaletteColor(role="neutral1", hex="#ECEBE4", usage_pct=15.0),
                PaletteColor(role="neutral2", hex="#D8D6CD", usage_pct=10.0),
                PaletteColor(role="card-bg", hex="#FFFFFF", usage_pct=15.0),
                PaletteColor(role="line", hex="#9A988F", usage_pct=10.0),
            ]
        ),
        fonts=FontPair(heading="Arial", body="Arial"),
        type_scale=TypeScale(
            entries=[
                TypeScaleEntry(name="body", size_pt=16, weight="regular"),
                TypeScaleEntry(name="small", size_pt=12, weight="regular"),
                TypeScaleEntry(name="caption", size_pt=9, weight="regular"),
            ]
        ),
        margins=Margins(),
        grid=Grid(columns=12),
        corner_radii=0.02,
        line_weight_pt=1.0,
    )


def _spec() -> DiagramSpec:
    return DiagramSpec(
        graph=GraphSpec(
            directed=True,
            nodes=[
                GraphNode(id="A", label="Start", kind="start"),
                GraphNode(id="B", label="Ingest"),
                GraphNode(id="C", label="Decide", kind="decision"),
                GraphNode(id="E", label="Branch"),
                GraphNode(id="F", label="Send"),
                GraphNode(id="G", label="Done", kind="end"),
            ],
            edges=[
                GraphEdge(src=s, dst=d)
                for s, d in [("A", "B"), ("B", "C"), ("A", "E"), ("E", "F"), ("F", "G"), ("C", "G")]
            ],
        )
    )


def _new_slide(P: Presentation):
    P.slide_width = Emu(SLIDE_W)
    P.slide_height = Emu(SLIDE_H)
    return P.slides.add_slide(P.slide_layouts[6])


def test_render_roundtrip(tmp_path):
    P = Presentation()
    slide = _new_slide(P)
    info = render_diagram(slide, _spec(), REGION, _style(), CANVAS)
    assert info.node_count == 6
    assert info.edge_count == 6
    assert not info.warnings

    path = tmp_path / "diagram.pptx"
    P.save(path)

    reopened = Presentation(path)
    shapes = list(reopened.slides[0].shapes)
    connectors = [s for s in shapes if isinstance(s, Connector)]
    # 6 edges, each elbow-routed with 4 points -> 3 chained straight connectors
    assert len(connectors) == 6 * 3
    assert len(shapes) == 6 + 6 * 3

    labels = {"Start", "Ingest", "Decide", "Branch", "Send", "Done"}
    text_shapes = [s for s in shapes if not isinstance(s, Connector)]
    assert {s.text_frame.text for s in text_shapes} == labels
    assert all(s.has_text_frame for s in text_shapes)

    connector = connectors[0]
    assert connector.line.color.rgb == RGBColor.from_string("9A988F")
    assert connector.line.width == Pt(1.0)


def test_mermaid_spec_renders_dashed_edge(tmp_path):
    P = Presentation()
    slide = _new_slide(P)
    spec = DiagramSpec(mermaid="flowchart LR\n  A --> B\n  A -.-> C")
    info = render_diagram(slide, spec, REGION, _style(), CANVAS)
    assert info.node_count == 3
    assert info.edge_count == 2
    assert not info.warnings
    P.save(tmp_path / "diagram_lr.pptx")

    connectors = [
        s for s in Presentation(tmp_path / "diagram_lr.pptx").slides[0].shapes
        if isinstance(s, Connector)
    ]
    dashed = [c for c in connectors if c.line.dash_style == MSO_LINE_DASH_STYLE.DASH]
    assert len(dashed) >= 3  # one dashed edge: 3 chained segments


def test_bad_spec_returns_warnings_not_exceptions():
    P = Presentation()
    slide = _new_slide(P)
    bad = DiagramSpec(mermaid="flowchart TD\n  A --o B")
    info = render_diagram(slide, bad, REGION, _style(), CANVAS)
    assert info.node_count == 0
    assert info.warnings


def test_empty_spec_returns_warnings_not_exceptions():
    P = Presentation()
    slide = _new_slide(P)
    info = render_diagram(slide, None, REGION, _style(), CANVAS)
    assert info.node_count == 0
    assert info.warnings
