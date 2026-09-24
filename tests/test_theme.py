"""Theme generation: StyleProfile from a brief + PowerPoint theme XML."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation

from deckforge_core.schemas.pack import StyleProfile
from deckforge_core.theme.contrast import contrast_ratio
from deckforge_core.theme.generator import palette_from_brief, style_from_brief
from deckforge_core.theme.theme_builder import THEME_XML_TEMPLATE, apply_theme

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _brief_style(seed: int = 5) -> StyleProfile:
    return style_from_brief("calm, editorial, warm neutrals", seed=seed)


def test_brief_builds_valid_profile():
    style = _brief_style()
    assert isinstance(style, StyleProfile)
    assert StyleProfile.model_validate(style.model_dump()) == style


def test_profile_palette_roles_present():
    style = _brief_style()
    for role in ("bg", "text", "muted-text", "accent1", "accent2", "neutral1", "line"):
        assert style.palette.has(role)


def test_contrast_guarantees_hold():
    style = _brief_style()
    assert contrast_ratio(style.palette.hex("text"), style.palette.hex("bg")) >= 4.5
    assert contrast_ratio(style.palette.hex("accent1"), style.palette.hex("bg")) >= 3.0
    assert contrast_ratio(style.palette.hex("muted-text"), style.palette.hex("bg")) >= 3.0


def test_deterministic_same_seed_same_dump():
    a = _brief_style(seed=5).model_dump(mode="json")
    b = _brief_style(seed=5).model_dump(mode="json")
    assert a == b


def test_seed_changes_palette_or_notes():
    a = _brief_style(seed=5)
    b = _brief_style(seed=6)
    a_pal = [c.hex for c in a.palette.colors]
    b_pal = [c.hex for c in b.palette.colors]
    assert a_pal != b_pal or a.notes != b.notes


def test_no_mood_keywords_yields_deterministic_default():
    a = style_from_brief("quarterly review meeting", seed=2)
    b = style_from_brief("quarterly review meeting", seed=2)
    assert isinstance(a, StyleProfile)
    assert a.model_dump() == b.model_dump()
    assert any("default" in n for n in a.notes)


def test_palette_from_brief_reuse():
    palette = palette_from_brief("energetic bold launch", seed=1)
    assert palette.has("accent1") and palette.has("bg")
    assert contrast_ratio(palette.hex("text"), palette.hex("bg")) >= 4.5


def _xml_root(style: StyleProfile, ns: dict[str, str]) -> etree._Element:
    return etree.fromstring(THEME_XML_TEMPLATE(style).encode("utf-8")).xpath(
        "/a:theme", namespaces=ns
    )[0]


def test_theme_xml_parses_with_lxml():
    style = _brief_style()
    root = etree.fromstring(THEME_XML_TEMPLATE(style).encode("utf-8"))
    assert root.tag == f"{{{_A_NS}}}theme"
    assert root.get("name") == "DeckForge Theme"


def test_clrscheme_has_exactly_twelve_children_in_order():
    style = _brief_style()
    ns = {"a": _A_NS}
    clr = _xml_root(style, ns).xpath("a:themeElements/a:clrScheme", namespaces=ns)[0]
    tags = [e.tag.split("}")[1] for e in clr]
    assert tags == [
        "dk1", "lt1", "dk2", "lt2",
        "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
        "hlink", "folHlink",
    ]


def test_clrscheme_carries_palette_hexes():
    style = _brief_style()
    ns = {"a": _A_NS}
    root = _xml_root(style, ns)
    srgb = root.xpath("//a:srgbClr/@val", namespaces=ns)
    blobish = "".join(srgb).lower()
    assert style.palette.hex("accent1")[1:].lower() in blobish
    assert style.palette.hex("accent2")[1:].lower() in blobish


def test_fontscheme_uses_profile_families():
    style = _brief_style()
    ns = {"a": _A_NS}
    root = _xml_root(style, ns)
    major = root.xpath("a:themeElements/a:fontScheme/a:majorFont/a:latin", namespaces=ns)[0]
    minor = root.xpath("a:themeElements/a:fontScheme/a:minorFont/a:latin", namespaces=ns)[0]
    assert major.get("typeface") == style.fonts.heading
    assert minor.get("typeface") == style.fonts.body


def test_fmtscheme_lists_have_three_children():
    style = _brief_style()
    ns = {"a": _A_NS}
    fmt = _xml_root(style, ns).xpath("a:themeElements/a:fmtScheme", namespaces=ns)[0]
    tags = [e.tag.split("}")[1] for e in fmt]
    assert tags == ["fillStyleLst", "lnStyleLst", "effectStyleLst", "bgFillStyleLst"]
    for child in fmt:
        assert len(child) == 3


def test_set_theme_rewrites_theme_part_and_reopens(tmp_path):
    style = _brief_style()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    warnings = apply_theme(prs, style)
    assert warnings == []

    out = tmp_path / "styled.pptx"
    prs.save(out)

    reopened = Presentation(str(out))
    theme_parts = [
        p for p in reopened.part.package.iter_parts() if str(p.partname).endswith("theme/theme1.xml")
    ]
    assert len(theme_parts) == 1
    blob = theme_parts[0].blob.decode("utf-8")
    assert style.palette.hex("accent1")[1:] in blob
    assert style.fonts.heading in blob
    assert style.fonts.body in blob
    assert "DeckForge Theme" in blob


def test_theme_xml_for_no_keyword_profile_parses():
    style = style_from_brief("no obvious mood here")
    root = etree.fromstring(THEME_XML_TEMPLATE(style).encode("utf-8"))
    assert root.tag == f"{{{_A_NS}}}theme"
