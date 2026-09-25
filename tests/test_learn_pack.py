"""End-to-end pack learning (Workstream B): learn_pack over a mixed corpus
(python-pptx synthetic + shipped corpus decks) must detect multiple
archetypes, derive a style profile, persist pack files and round-trip."""

from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from deckforge_core.analysis import analyse_decks, learn_pack
from deckforge_core.ingest import extract_directory
from deckforge_core.renderer.pack_io import load_pack

REPO_ROOT = Path(__file__).resolve().parents[1]
CORPUS = REPO_ROOT / "corpus" / "raw"

HEADING_FONT = "Georgia"
BODY_FONT = "Arial"


def _brightness(hex_col: str) -> float:
    h = hex_col.lstrip("#")
    r, g, b = (int(h[i : i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b


def _add_text(slide, x_in, y_in, w_in, h_in, lines, *, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(
        Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in)
    )
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, size, bold in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = HEADING_FONT if bold else BODY_FONT
    return box


def _png_bytes(color: tuple[int, int, int]) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (200, 150), color).save(buf, "PNG")
    return buf.getvalue()


def _rich_deck(path: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    _add_text(s, 0.4, 0.4, 9.2, 0.4, [("Q3 2026", 12, True)])
    _add_text(s, 0.4, 1.2, 9.2, 0.8, [("Building the category", 30, True)])
    _add_text(s, 0.4, 2.3, 9.2, 0.5, [("A working document for the team", 16, False)])

    s = prs.slides.add_slide(blank)
    _add_text(s, 0.4, 2.0, 9.2, 1.8, [("91", 60, True)])
    _add_text(s, 0.4, 4.2, 9.2, 0.6, [("Users up year over year", 18, False)])

    s = prs.slides.add_slide(blank)
    _add_text(s, 0.4, 0.4, 9.2, 0.6, [("Where we win", 20, True)])
    _add_text(
        s, 0.4, 2.0, 2.8, 2.2,
        [("Bookings", 18, True), ("Bookings grew every single week this quarter.", 14, False)],
        align=PP_ALIGN.CENTER,
    )
    _add_text(
        s, 3.7, 2.0, 2.8, 2.2,
        [("Support", 18, True), ("Support answered every question within the hour.", 14, False)],
        align=PP_ALIGN.CENTER,
    )
    _add_text(
        s, 7.0, 2.0, 2.8, 2.2,
        [("Retention", 18, True), ("Retention stayed at ninety five percent this quarter.", 14, False)],
        align=PP_ALIGN.CENTER,
    )

    s = prs.slides.add_slide(blank)
    _add_text(s, 0.4, 0.4, 9.2, 0.6, [("Revenue trend", 20, True)])
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3"]
    chart_data.add_series("Revenue", (1.0, 2.0, 3.0))
    s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.4), Inches(1.8), Inches(9.2), Inches(4.6),
        chart_data,
    )

    s = prs.slides.add_slide(blank)
    table_shape = s.shapes.add_table(3, 3, Inches(0.4), Inches(1.6), Inches(6.0), Inches(3.0))
    rows = [
        ["Metric", "Q1", "Q2"],
        ["Revenue", "1.0m", "2.0m"],
        ["Customers", "400", "650"],
    ]
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            table_shape.table.cell(ri, ci).text = cell

    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(io.BytesIO(_png_bytes((150, 90, 60))), Inches(0.4), Inches(1.5), width=Inches(4.0))
    _add_text(s, 5.0, 1.4, 4.6, 0.6, [("Customers love it", 30, True)])
    _add_text(
        s, 5.0, 2.4, 4.6, 2.6,
        [
            ("Teams that ship weekly see retention rise steadily across the year.", 16, False),
            ("The feedback loop keeps shortening with every single quarter of work.", 16, False),
        ],
    )

    prs.save(path)


@pytest.fixture(scope="module")
def input_dir(tmp_path_factory) -> Path:
    inputs = sorted(CORPUS.glob("*.pptx"))
    if not inputs:
        pytest.skip("corpus/raw has no decks; cannot run learn_pack end-to-end")
    folder = tmp_path_factory.mktemp("input")
    for deck in inputs[:4]:
        shutil.copy2(deck, folder / deck.name)
    _rich_deck(folder / "rich-editorial.pptx")
    return folder


def test_learn_pack_detects_multiple_archetypes(input_dir: Path, tmp_path: Path):
    pack = learn_pack(input_dir, "mixed", target_dir=tmp_path / "out", use_cache=True)
    assert pack.name == "mixed"
    assert pack.source_deck_count == 5
    covered = set(pack.archetypes)
    assert len(covered) >= 4
    assert "title" in covered
    assert "section-divider" in covered


def test_learn_pack_writes_files_and_round_trips(input_dir: Path, tmp_path: Path):
    pack = learn_pack(input_dir, "mixed", target_dir=tmp_path / "out", use_cache=True)
    assert (tmp_path / "out" / "pack.json").exists()
    assert (tmp_path / "out" / "style_profile.json").exists()
    blueprints = list((tmp_path / "out" / "blueprints").glob("*.json"))
    assert blueprints
    loaded, library = load_pack(tmp_path / "out")
    assert loaded.style.model_dump(mode="json") == pack.style.model_dump(mode="json")
    assert library.archetypes()


def test_style_palette_has_full_role_set(input_dir: Path, tmp_path: Path):
    pack = learn_pack(input_dir, "mixed", target_dir=tmp_path / "out", use_cache=True)
    roles = [c.role for c in pack.style.palette.colors]
    assert len(roles) >= 5
    text_hex = pack.style.palette.hex("text")
    bg_hex = pack.style.palette.hex("bg")
    assert _brightness(text_hex) < _brightness(bg_hex)
    assert pack.style.fonts.heading
    assert pack.style.fonts.body


def test_learn_pack_is_deterministic(input_dir: Path, tmp_path: Path):
    first = learn_pack(input_dir, "mixed", target_dir=tmp_path / "a", use_cache=True)
    second = learn_pack(input_dir, "mixed", target_dir=tmp_path / "b", use_cache=True)
    assert first.style.model_dump(mode="json") == second.style.model_dump(mode="json")
    assert first.archetypes == second.archetypes


def test_analyse_decks_reports_counts(input_dir: Path):
    result = extract_directory(input_dir, use_cache=True)
    report = analyse_decks(result.decks)
    assert report.deck_count == 5
    assert report.slide_count > 10
    assert set(report.archetype_counts) >= {"title", "section-divider"}
    assert isinstance(report.new_archetypes, list)
    assert report.aspect_ratios
