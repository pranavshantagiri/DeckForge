"""Widget layer of the DeckForge GUI, driven headless.

``QT_QPA_PLATFORM=offscreen`` must be set before Qt is imported, so these tests
run without a display and without a real event loop. Live PowerPoint COM is
never used: previews go through a fake renderer and the "none" backend otherwise.
"""

from __future__ import annotations

import os
import shutil
import time
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.pop("DECKFORGE_TEST_COM_RENDER", None)

pytest.importorskip("PySide6", reason="the gui extra (PySide6) is not installed")

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from deckforge_core.config import Settings  # noqa: E402
from deckforge_core.gui import models  # noqa: E402
from deckforge_core.gui.generate_page import GeneratePage  # noqa: E402
from deckforge_core.gui.main import MainWindow, build_app  # noqa: E402
from deckforge_core.gui.packs_page import LearnPackDialog, PacksPage  # noqa: E402
from deckforge_core.gui.settings_page import SettingsPage  # noqa: E402
from deckforge_core.providers import NoneRenderer  # noqa: E402
from deckforge_core.providers.render import RenderBackend, RenderedSlide  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[1]
CORPUS = REPO_ROOT / "corpus" / "raw"


class FakeSecretStore:
    """In-memory stand-in for the OS credential store."""

    SERVICE = "DeckForge"

    def __init__(self) -> None:
        self.values: dict[str, str] = {}

    def get(self, provider: str, default: str | None = None) -> str | None:
        return self.values.get(provider, default)

    def set(self, provider: str, value: str) -> None:
        self.values[provider] = value

    def delete(self, provider: str) -> None:
        self.values.pop(provider, None)


class FakeRenderer:
    backend = RenderBackend.POWERPOINT_COM

    def render(self, pptx_path, out_dir, *, slides=None, dpi=96):
        target = Path(out_dir)
        target.mkdir(parents=True, exist_ok=True)
        count = len(Presentation(str(pptx_path)).slides)
        results = []
        for number in range(1, count + 1):
            png = target / f"slide_{number:03d}.png"
            Image.new("RGB", (320, 180), (30 * number % 255, 90, 140)).save(png, "PNG")
            results.append(
                RenderedSlide(number=number, image_path=str(png), width_px=320, height_px=180)
            )
        return results


class FakeAutodetect:
    def __init__(self) -> None:
        self.com = True
        self.libreoffice = False

    def resolve(self, preferred: str = "auto"):
        if preferred == RenderBackend.NONE.value:
            return NoneRenderer()
        return FakeRenderer()


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture(scope="session")
def source_folder(tmp_path_factory) -> Path:
    decks = sorted(CORPUS.glob("*.pptx"))[:2]
    if not decks:
        pytest.skip("corpus/raw has no decks; cannot learn a pack")
    folder = tmp_path_factory.mktemp("gui-source-decks")
    for deck in decks:
        shutil.copy2(deck, folder / deck.name)
    return folder


@pytest.fixture(scope="session")
def packs_root(tmp_path_factory, source_folder) -> Path:
    root = tmp_path_factory.mktemp("gui-packs")
    models.learn_pack_from_folder(
        source_folder, "gui-pack", target_dir=root / "gui-pack", use_cache=False
    )
    return root


@pytest.fixture
def settings(tmp_path) -> Settings:
    instance = Settings(root=tmp_path)
    instance.update(llm_provider="mock")
    return instance


@pytest.fixture
def window(qapp, settings, packs_root, tmp_path) -> MainWindow:
    win = MainWindow(
        settings=settings,
        secrets=FakeSecretStore(),
        packs_dir=packs_root,
        decks_dir=tmp_path / "decks",
    )
    # previews stay off: no test may reach PowerPoint COM or LibreOffice
    win.generate_page.backend_combo.setCurrentIndex(
        win.generate_page.backend_combo.findData("none")
    )
    yield win
    win.deleteLater()


def _wait_until(predicate, timeout: float = 60.0) -> bool:
    """Pump the event loop until ``predicate`` holds; no ``exec()`` is started."""
    app = QApplication.instance()
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return True
        time.sleep(0.01)
    return False


# --------------------------------------------------------------------------- #
# Window shell
# --------------------------------------------------------------------------- #
def test_build_app_creates_themed_window(qapp):
    app, win = build_app([])
    assert app is qapp
    assert win.tabs.count() == 3
    assert [win.tabs.tabText(i) for i in range(3)] == ["Packs", "Generate", "Settings"]
    assert app.styleSheet()
    assert win.statusBar().currentMessage()
    win.deleteLater()


def test_main_window_lists_packs_and_navigates(window, packs_root):
    assert window.packs_page.list.count() == 1
    assert window.generate_page.pack_combo.count() == 1
    assert window.generate_page.pack_combo.currentText() == "gui-pack"
    window.show_tab(1)
    assert window.tabs.currentIndex() == 1
    window.refresh_all()
    assert window.packs_page.list.count() == 1


def test_pack_chosen_switches_to_generate(window):
    window.packs_page.select("gui-pack")
    window.packs_page._use_pack()
    assert window.tabs.currentIndex() == 1
    assert window.generate_page.pack_combo.currentText() == "gui-pack"
    assert "gui-pack" in window.statusBar().currentMessage()


def test_close_is_ignored_while_a_task_runs(window):
    assert window.packs_page.runner.start(lambda progress=None: time.sleep(1.5))
    assert window.busy
    assert window.close() is False
    assert window.packs_page.runner.busy
    assert "Still working" in window.statusBar().currentMessage()
    assert _wait_until(lambda: not window.busy)
    assert window.close() is True


def test_packs_page_refresh_reports_empty_folder(qapp, tmp_path):
    page = PacksPage(tmp_path / "empty-packs")
    assert page.list.count() == 0
    assert "No packs" in page.status.text()
    assert not page.use_button.isEnabled()
    page.deleteLater()


# --------------------------------------------------------------------------- #
# Worker plumbing
# --------------------------------------------------------------------------- #
def test_task_passes_progress_keyword_and_returns(qapp):
    from deckforge_core.gui.workers import TaskRunner

    def work(value, *, progress):
        progress(40, "halfway")
        return value * 2

    runner = TaskRunner()
    progress: list[tuple[int, str]] = []
    results: list[object] = []
    done: list[int] = []
    assert runner.start(
        work,
        21,
        on_progress=lambda percent, message: progress.append((percent, message)),
        on_success=results.append,
        on_finished=lambda: done.append(1),
    )
    assert _wait_until(lambda: done)
    assert results == [42]
    assert progress == [(40, "halfway")]
    assert not runner.busy


def test_task_reports_exceptions_instead_of_crashing(qapp):
    from deckforge_core.gui.workers import TaskRunner

    def boom(*, progress):
        raise ValueError("kaboom")

    runner = TaskRunner()
    errors: list[str] = []
    done: list[int] = []
    assert runner.start(boom, on_error=errors.append, on_finished=lambda: done.append(1))
    assert _wait_until(lambda: done)
    assert errors == ["kaboom"]


def test_task_runner_refuses_a_second_task(qapp):
    from deckforge_core.gui.workers import TaskRunner

    runner = TaskRunner()
    done: list[int] = []
    assert runner.start(
        lambda progress=None: time.sleep(0.6), on_finished=lambda: done.append(1)
    )
    assert not runner.start(lambda progress=None: time.sleep(0.6))
    assert _wait_until(lambda: done)
    assert runner.busy is False


# --------------------------------------------------------------------------- #
# Packs page
# --------------------------------------------------------------------------- #
def test_packs_page_lists_learned_pack(window, packs_root):
    page = window.packs_page
    entry = page.selected_entry()
    assert entry is not None
    assert entry.name == "gui-pack"
    assert entry.deck_count == 2
    assert "decks" in entry.summary
    assert page.use_button.isEnabled()
    assert page.select("gui-pack")
    assert page.selected_name() == "gui-pack"
    assert not page.select("missing-pack")


def test_packs_page_learns_a_new_pack(qapp, tmp_path, source_folder):
    page = PacksPage(tmp_path / "packs")
    messages: list[str] = []
    result = page.learn(
        source_folder,
        "second pack",
        progress=lambda percent, message: messages.append(message),
    )
    assert result.pack.name == "second pack"
    assert (result.pack_dir / "pack.json").is_file()
    assert any("Reading" in message for message in messages)
    page.refresh()
    names = [page.list.item(i).text().splitlines()[0] for i in range(page.list.count())]
    assert names == ["second-pack"]
    assert "second pack" in page.list.item(0).text()
    page.deleteLater()


def test_packs_page_learn_runs_on_a_worker_thread(qapp, tmp_path, source_folder):
    page = PacksPage(tmp_path / "packs")
    errors: list[str] = []
    page.show_error = errors.append
    assert page.start_learn(source_folder, "threaded pack")
    assert _wait_until(
        lambda: not page.runner.busy and page.status.text().startswith("Learned")
    )
    assert errors == []
    assert page.list.count() == 1
    assert page.progress.value() == 100
    page.deleteLater()


def test_learned_pack_reaches_the_generate_page(window, source_folder):
    page = window.packs_page
    page.show_error = lambda message: pytest.fail(message)
    assert page.start_learn(source_folder, "third pack")
    assert _wait_until(
        lambda: not page.runner.busy and page.status.text().startswith("Learned")
    )
    combo = window.generate_page.pack_combo
    assert combo.count() == 2
    assert "third-pack" in [combo.itemText(i) for i in range(combo.count())]


def test_learn_pack_dialog_reads_its_fields(qapp, tmp_path):
    dialog = LearnPackDialog()
    dialog.folder_edit.setText(str(tmp_path))
    dialog.name_edit.setText("My Pack")
    folder, name = dialog.values()
    assert folder == tmp_path
    assert name == "My Pack"
    dialog.folder_edit.setText(str(tmp_path))
    dialog.name_edit.setText("")
    assert dialog.values()[1] == tmp_path.name
    dialog.deleteLater()


def test_packs_page_imports_a_dfpack(qapp, packs_root, tmp_path):
    from deckforge_core.renderer.pack_io import export_dfpack

    archive = tmp_path / "gui-pack.dfpack"
    export_dfpack(packs_root / "gui-pack", archive)
    destination = tmp_path / "packs"
    page = PacksPage(destination)
    assert page.list.count() == 0
    entry = page.import_pack(archive)
    assert entry.name == "gui-pack"
    page.refresh()
    assert page.list.count() == 1
    page.deleteLater()


# --------------------------------------------------------------------------- #
# Generate page
# --------------------------------------------------------------------------- #
def test_generate_page_builds_request_from_fields(window, packs_root, tmp_path):
    page = window.generate_page
    page.brief_edit.setPlainText("Quarterly review for the board")
    page.audience_edit.setText("board members")
    page.slide_count.setValue(7)
    page.tone_combo.setCurrentText("executive")
    page.aspect_combo.setCurrentIndex(page.aspect_combo.findData("4:3"))
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    request = page.build_request()
    assert request.pack_dir == packs_root / "gui-pack"
    assert request.brief == "Quarterly review for the board"
    assert request.audience == "board members"
    assert request.slide_count == 7
    assert request.tone == "executive"
    assert request.aspect_ratio == "4:3"
    assert request.render_backend == "none"
    assert request.use_llm is False
    assert request.options().slide_count == 7


def test_generate_page_needs_a_pack(qapp, tmp_path, settings):
    page = GeneratePage(tmp_path / "no-packs", settings)
    with pytest.raises(ValueError, match="select a format pack"):
        page.build_request()
    errors: list[str] = []
    page.show_error = errors.append
    page.on_generate_clicked()
    assert errors and "select a format pack" in errors[0]
    assert page.status.text().startswith("Failed")
    page.deleteLater()


def test_generate_page_renders_and_previews(window):
    page = window.generate_page
    page.brief_edit.setPlainText("Q3 results for the leadership team")
    page.slide_count.setValue(5)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    result = page.run_generation()
    page.show_result(result)

    assert result.pptx_path.is_file()
    assert result.slide_count == 5
    assert page.thumb_list.count() == 5
    assert page.open_button.isEnabled()
    assert page.reveal_button.isEnabled()
    assert page.preview_button.isEnabled()
    assert "5 slides" in page.status.text()
    assert page.issues_list.count() >= 1
    assert page.current_pptx() == result.pptx_path
    presentation = Presentation(str(result.pptx_path))
    assert len(presentation.slides) == 5
    assert result.pptx_path.parent.parent == window.generate_page.decks_dir


def test_generate_page_thumbnail_stage_switches(window, monkeypatch):
    monkeypatch.setattr(models, "RendererAutodetect", FakeAutodetect)
    page = window.generate_page
    page.brief_edit.setPlainText("Preview me")
    page.slide_count.setValue(4)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("auto"))
    result = page.run_generation()
    page.show_result(result)
    page.stage.resize(640, 360)
    assert result.rendered_previews == 4
    page.thumb_list.setCurrentRow(2)
    assert not page.stage.pixmap().isNull()
    assert "slide_003.png" in page.stage.toolTip()
    page.thumb_list.setCurrentRow(0)
    assert page.thumb_list.item(0).icon().isNull() is False


def test_generate_page_without_images_shows_a_text_card(window):
    page = window.generate_page
    page.brief_edit.setPlainText("No preview machine")
    page.slide_count.setValue(3)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    result = page.run_generation()
    page.show_result(result)
    page.thumb_list.setCurrentRow(1)
    assert page.stage.pixmap().isNull()
    assert "No preview image" in page.stage.text()
    assert "no previews" in page.status.text()


def test_generate_page_opens_the_deck_with_the_default_app(window, monkeypatch):
    opened: list[str] = []
    monkeypatch.setattr(models, "open_in_default_app", lambda path: opened.append(str(path)))
    page = window.generate_page
    page.brief_edit.setPlainText("Open me")
    page.slide_count.setValue(3)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    page.show_result(page.run_generation())
    assert page.open_button.isEnabled()
    page.on_open_clicked()
    assert opened == [str(page.current_pptx())]
    assert "Opened" in page.status.text()


def test_generate_page_repreview_uses_the_fake_renderer(window, monkeypatch):
    monkeypatch.setattr(models, "RendererAutodetect", FakeAutodetect)
    page = window.generate_page
    page.brief_edit.setPlainText("Repreview me")
    page.slide_count.setValue(3)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    page.show_result(page.run_generation())
    assert page._result.rendered_previews == 0
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("auto"))
    page.on_preview_clicked()
    assert page._result.rendered_previews == 3
    assert page.thumb_list.item(0).icon().isNull() is False


def test_generate_page_start_generation_uses_the_runner(window):
    page = window.generate_page
    page.brief_edit.setPlainText("Threaded deck")
    page.slide_count.setValue(4)
    page.backend_combo.setCurrentIndex(page.backend_combo.findData("none"))
    assert page.start_generation()
    assert _wait_until(lambda: not page.runner.busy and page.thumb_list.count() == 4)
    assert page.current_pptx().is_file()
    assert page.progress.value() == 100


def test_generate_page_reports_failures(window, monkeypatch):
    page = window.generate_page
    errors: list[str] = []
    page.show_error = errors.append

    def _boom(*args, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(models, "generate", _boom)
    page.brief_edit.setPlainText("Will fail")
    assert page.start_generation()
    assert _wait_until(lambda: not page.runner.busy)
    assert errors == ["boom"]
    assert page.status.text() == "Failed: boom"


# --------------------------------------------------------------------------- #
# Settings page
# --------------------------------------------------------------------------- #
def test_settings_page_saves_non_secret_preferences(qapp, tmp_path):
    settings = Settings(root=tmp_path)
    page = SettingsPage(settings, FakeSecretStore())
    page.local_only.setChecked(False)
    page.provider_combo.setCurrentIndex(page.provider_combo.findData("mock"))
    page.model_plan.setText("test-model")
    page.max_iterations.setValue(2)
    page.engine_combo.setCurrentIndex(page.engine_combo.findData("none"))
    saved: list[int] = []
    page.settings_saved.connect(lambda: saved.append(1))
    page.save()
    assert saved == [1]
    assert settings.get("local_only") is False
    assert settings.get("llm_model_plan") == "test-model"
    assert settings.get("max_qa_iterations") == 2
    assert settings.get("render_engine") == "none"
    assert "Settings saved." in page.status.text()
    assert page.storage_path() == tmp_path / Settings.FILE_NAME
    page.deleteLater()


def test_settings_page_keeps_keys_out_of_the_settings_file(qapp, tmp_path):
    settings = Settings(root=tmp_path)
    secrets = FakeSecretStore()
    page = SettingsPage(settings, secrets)
    assert not page.has_key("anthropic")
    assert "not set" in page.key_edits["anthropic"].placeholderText()
    page.key_edits["anthropic"].setText("sk-ant-super-secret")
    page.save_key("anthropic")
    assert secrets.get("anthropic") == "sk-ant-super-secret"
    assert page.has_key("anthropic")
    assert page.key_edits["anthropic"].text() == ""
    assert "stored" in page.key_edits["anthropic"].placeholderText()
    page.save()
    stored = (tmp_path / Settings.FILE_NAME).read_text(encoding="utf-8")
    assert "sk-ant-super-secret" not in stored
    assert "anthropic" in stored  # the provider name is fine, the key is not
    page.delete_key("anthropic")
    assert not page.has_key("anthropic")
    assert "Removed" in page.status.text()
    page.deleteLater()


def test_settings_page_masks_and_reveals_keys(qapp, tmp_path):
    page = SettingsPage(Settings(root=tmp_path), FakeSecretStore())
    field = page.key_edits["openai"]
    assert field.echoMode() == field.EchoMode.Password
    page.reveal_button.setChecked(True)
    assert field.echoMode() == field.EchoMode.Normal
    page.reveal_button.setChecked(False)
    assert field.echoMode() == field.EchoMode.Password
    page.deleteLater()


def test_settings_page_ignores_empty_keys(qapp, tmp_path):
    secrets = FakeSecretStore()
    page = SettingsPage(Settings(root=tmp_path), secrets)
    page.save_key("gemini")
    assert secrets.values == {}
    assert "No key entered" in page.status.text()
    page.deleteLater()


def test_settings_page_shows_backend_hint(qapp, tmp_path):
    page = SettingsPage(Settings(root=tmp_path), FakeSecretStore())
    assert "Slide renderers detected" in page.backend_hint.text()
    page.deleteLater()
