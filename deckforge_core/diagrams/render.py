"""Render a DiagramSpec onto a python-pptx slide as native shapes and links.

Pure offline: no Office application and no Graphviz are ever touched. Node
bodies are real, editable AutoShapes with solid fills; edges are native
``cxnSp`` connectors with explicit strokes. Multi-bend (elbow) routes are drawn
as a chain of straight connectors — one per consecutive pair of route points —
so any number of bend points works with the native connector API; a directed
edge ends with a small triangle arrowhead on its final segment.

Rendering never raises on a bad spec: parse/layout problems are collected into
:class:`DiagramInfo` warnings and the function draws whatever it still can.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from deckforge_core.diagrams import DiagramError
from deckforge_core.diagrams.layout import layout_graph
from deckforge_core.diagrams.mermaid import parse_mermaid
from deckforge_core.schemas.blueprints import SlotRegion
from deckforge_core.schemas.deck_plan import DiagramSpec, GraphSpec
from deckforge_core.schemas.pack import ColorPalette, StyleProfile


@dataclass
class DiagramInfo:
    """Summary of what was actually drawn on the slide."""

    node_count: int = 0
    edge_count: int = 0
    node_ids: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.warnings


def render_diagram(
    slide,
    spec: DiagramSpec | None,
    region: SlotRegion,
    style: StyleProfile,
    canvas_emu: tuple[int, int],
    *,
    orientation: str = "TD",
) -> DiagramInfo:
    """Draw ``spec``'s graph inside ``region`` on ``slide``.

    ``canvas_emu`` is the slide's (width, height) in EMU; relative coordinates
    from the layout are scaled against it. Returns a :class:`DiagramInfo`.
    """
    warnings: list[str] = []
    graph, mermaid_orientation, resolution = _effective_graph(spec)
    warnings.extend(resolution)
    if mermaid_orientation:
        orientation = mermaid_orientation
    if not graph.nodes:
        warnings.append("diagram has no nodes; nothing rendered")

    try:
        layout = layout_graph(graph, region, directed=graph.directed, orientation=orientation)
    except DiagramError as exc:
        warnings.append(str(exc))
        return DiagramInfo(warnings=warnings)
    warnings.extend(layout.warnings)

    cw, ch = canvas_emu
    palette = style.palette
    fill_hex = _role_hex(palette, ("card-bg", "neutral1", "bg"), default="#FFFFFF")
    line_hex = _role_hex(palette, ("line", "neutral2", "neutral1", "muted-text"), default="#9A988F")
    text_hex = _role_hex(palette, ("text", "muted-text"), default="#1C1C1A")
    label_hex = _role_hex(palette, ("muted-text", "line", "text"), default="#6B6B66")
    weight = style.line_weight_pt if style.line_weight_pt > 0 else 0.75
    font_name = style.fonts.body
    node_pt = _scale_pt(style, ("small", "body"), 12.0)
    label_pt = _scale_pt(style, ("caption", "small"), 9.0)

    shapes = slide.shapes
    for pn in layout.nodes:
        left = int(pn.x * cw)
        top = int(pn.y * ch)
        width = max(int(pn.w * cw), 1)
        height = max(int(pn.h * ch), 1)
        shp = shapes.add_shape(_shape_kind(pn.kind), left, top, width, height)
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill_hex)
        shp.line.color.rgb = _rgb(line_hex)
        shp.line.width = Pt(weight)
        shp.shadow.inherit = False
        if pn.kind == "process":
            shp.adjustments[0] = _rounded_adjustment(style, pn.w, pn.h)
        _set_node_text(shp, pn.label or pn.id, font_name, node_pt, _rgb(text_hex))

    for pe in layout.edges:
        pts = [(int(x * cw), int(y * ch)) for x, y in pe.route]
        segments = []
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            cxn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
            cxn.line.color.rgb = _rgb(line_hex)
            cxn.line.width = Pt(weight)
            if pe.style == "dashed":
                cxn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            segments.append(cxn)
        if pe.directed and segments:
            _add_arrow_head(segments[-1])
        if pe.label:
            _render_edge_label(
                slide, pe.label, pe.route, label_hex, font_name, label_pt, cw, ch
            )

    node_ids = [pn.id for pn in layout.nodes]
    return DiagramInfo(
        node_count=len(node_ids),
        edge_count=len(layout.edges),
        node_ids=node_ids,
        warnings=warnings,
    )


# --------------------------------------------------------------------------- #
# Spec resolution
# --------------------------------------------------------------------------- #
def _effective_graph(spec: DiagramSpec | None) -> tuple[GraphSpec, str | None, list[str]]:
    """Resolve the source graph + orientation; parse errors become warnings."""
    if spec is None:
        return GraphSpec(), None, ["no diagram spec supplied"]
    if spec.graph is not None:
        return spec.graph, None, []
    if spec.mermaid:
        try:
            graph, orientation = parse_mermaid(spec.mermaid)
            return graph, orientation, []
        except DiagramError as exc:
            return GraphSpec(), None, [f"mermaid parse failed: {exc}"]
    return GraphSpec(), None, ["diagram spec has neither graph nor mermaid source"]


# --------------------------------------------------------------------------- #
# Palette + typography helpers
# --------------------------------------------------------------------------- #
def _role_hex(palette: ColorPalette, roles: tuple[str, ...], *, default: str) -> str:
    """First palette role that exists, else dominant, else ``default``."""
    for role in roles:
        if role and palette.has(role):
            return palette.hex(role)
    try:
        return palette.dominant()
    except ValueError:
        return default


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _scale_pt(style: StyleProfile, names: tuple[str, ...], default: float) -> float:
    for name in names:
        try:
            return style.type_scale.entry(name).size_pt
        except KeyError:
            continue
    return default


# --------------------------------------------------------------------------- #
# Shape + connector builders
# --------------------------------------------------------------------------- #
def _shape_kind(kind: str):
    if kind == "decision":
        return MSO_SHAPE.DIAMOND
    if kind in ("start", "end"):
        return MSO_SHAPE.OVAL
    return MSO_SHAPE.ROUNDED_RECTANGLE


def _rounded_adjustment(style: StyleProfile, w: float, h: float) -> float:
    """Map the pack's relative corner radius onto the shape's 0..0.5 adjustment."""
    if style.corner_radii <= 0:
        return 0.0
    frac = style.corner_radii / max(min(w, h), 1e-9)
    return max(0.06, min(0.15, frac))


def _set_node_text(shp, text: str, font_name: str, size_pt: float, color: RGBColor) -> None:
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color


def _add_arrow_head(connector) -> None:
    """Set a triangle tail-end arrow on the connector's line element."""
    ln = connector.line._get_or_add_ln()
    tail = ln.find(qn("a:tailEnd"))
    if tail is None:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _render_edge_label(
    slide, label: str, route: list[tuple[float, float]], hex_color: str, font_name: str,
    size_pt: float, cw: int, ch: int,
) -> None:
    if len(route) >= 4:
        mx = (route[1][0] + route[2][0]) / 2
        my = (route[1][1] + route[2][1]) / 2
    else:
        mx = (route[0][0] + route[-1][0]) / 2
        my = (route[0][1] + route[-1][1]) / 2
    left, top = int(mx * cw), int(my * ch)
    tb = slide.shapes.add_textbox(left, top, int(0.1 * cw), int(0.03 * ch))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _rgb(hex_color)
