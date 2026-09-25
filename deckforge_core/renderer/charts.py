"""Native python-pptx charts styled strictly from the pack's StyleProfile.

Only flat (non-3D, non-shadow) chart types are emitted. Series fills rotate
through the palette accents; axes use the body font at small size; gridlines
are neutral; the legend follows the spec.
"""

from __future__ import annotations

from datetime import datetime, timezone
from typing import List

import xlsxwriter.workbook as _xlsxworkbook
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt

from deckforge_core.schemas.deck_plan import ChartSpec
from deckforge_core.schemas.pack import FormatPack

# Pin the creation timestamp python-pptx's xlsxwriter workbook embeds in
# ``docProps/core.xml`` so two charts with identical data produce a
# byte-identical embedded workbook (reproducible builds).
_FIXED_CREATE_TIME = datetime(1980, 1, 1, 0, 0, 0, tzinfo=timezone.utc)
_ORIGINAL_WORKBOOK_INIT = _xlsxworkbook.Workbook.__init__


def _pinned_workbook_init(self, *args, **kwargs):
    _ORIGINAL_WORKBOOK_INIT(self, *args, **kwargs)
    self.createtime = _FIXED_CREATE_TIME
    self.set_properties({"created": _FIXED_CREATE_TIME})


_xlsxworkbook.Workbook.__init__ = _pinned_workbook_init

CHART_TYPE_MAP: dict[str, XL_CHART_TYPE] = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
}

_DEFAULT_ROTATE: tuple[str, ...] = ("accent1", "accent2", "neutral1", "text")


def _hex(pack: FormatPack, *roles: str) -> str:
    palette = pack.style.palette
    for role in roles:
        try:
            return palette.hex(role)
        except KeyError:
            continue
    return palette.colors[0].hex


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#"))


def _series_roles(spec: ChartSpec, count: int) -> List[str]:
    roles: List[str] = []
    for i in range(count):
        if i < len(spec.series) and spec.series[i].color_role:
            roles.append(spec.series[i].color_role)
        else:
            roles.append(_DEFAULT_ROTATE[i % len(_DEFAULT_ROTATE)])
    return roles


def add_chart(
    slide,
    spec: ChartSpec,
    left: int,
    top: int,
    width: int,
    height: int,
    pack: FormatPack,
) -> None:
    """Draw ``spec`` as a native chart on ``slide``. Mutates the slide only."""
    chart_type = CHART_TYPE_MAP.get((spec.chart_type or "column").lower())
    if chart_type is None:
        chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED

    chart_data = CategoryChartData()
    chart_data.categories = list(spec.categories)
    for series in spec.series:
        chart_data.add_series(series.name, tuple(series.values))

    graphic_frame = slide.shapes.add_chart(
        chart_type, left, top, width, height, chart_data
    )
    chart = graphic_frame.chart

    style = pack.style
    body_font = style.fonts.body
    small_pt = _small_pt(pack)
    is_pie = chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)

    series = list(chart.plots[0].series)
    roles = _series_roles(spec, len(series))

    if is_pie:
        for index, ser in enumerate(series):
            ser_role = roles[index] if index < len(roles) else "accent1"
            for point in ser.points:
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(_hex(pack, ser_role))
    else:
        for index, ser in enumerate(series):
            ser_role = roles[index] if index < len(roles) else "accent1"
            color = _rgb(_hex(pack, ser_role))
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = color
            if chart_type == XL_CHART_TYPE.LINE:
                ser.format.line.color.rgb = color

    # axes + gridlines (not applicable to pie/donut)
    if not is_pie:
        value_axis = chart.value_axis
        category_axis = chart.category_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = _rgb(
            _hex(pack, "muted-text", "neutral1", "text")
        )
        value_axis.major_gridlines.format.line.width = Pt(0.75)
        for axis in (value_axis, category_axis):
            axis.tick_labels.font.name = body_font
            axis.tick_labels.font.size = Pt(small_pt)

    # legend / data labels
    chart.has_legend = bool(spec.show_legend)
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if chart.plots:
        chart.plots[0].has_data_labels = bool(spec.data_labels)


def _small_pt(pack: FormatPack) -> float:
    for name in ("small", "caption", "body"):
        try:
            return pack.style.type_scale.entry(name).size_pt
        except KeyError:
            continue
    return 12.0
