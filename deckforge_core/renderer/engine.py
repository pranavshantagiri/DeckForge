"""Main render loop: Deck Plan + Format Pack -> native .pptx.

Renders every :class:`SlidePlan` onto a blank layout using the blueprint
geometry and the pack's style only. Never raises on missing content or style
entries: it warns into the caller-provided list and keeps going.
"""

from __future__ import annotations

import warnings as _warnings_mod
import zipfile
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge_core.renderer.canvas import Canvas, canvas_size
from deckforge_core.renderer.charts import add_chart
from deckforge_core.renderer.layout import PlacedSlot, layout_slots
from deckforge_core.renderer.tables import add_table
from deckforge_core.renderer.textfit import fit_paragraphs
from deckforge_core.schemas.blueprints import (
    BlueprintLibrary,
    ContentKind,
)
from deckforge_core.schemas.deck_plan import DeckPlan, SlidePlan, SlotValue
from deckforge_core.schemas.pack import FormatPack, TypeScaleEntry

_HEADING_STYLES = {"h1", "h2", "kicker", "big-number"}
_BULLET_CHARS = ("•", "-", "–", "▪", "●", "*", "‣")


class RenderResult:
    """Outcome of a render: where the file went, how many slides, warnings."""

    def __init__(
        self, out_path: Path, slide_count: int, warnings: List[str], plan_snapshot: str
    ):
        self.out_path = Path(out_path)
        self.slide_count = slide_count
        self.warnings = warnings
        self.plan_snapshot = plan_snapshot

    def __repr__(self) -> str:  # pragma: no cover - debug helper
        return (
            f"RenderResult(out_path={self.out_path!r}, slide_count={self.slide_count}, "
            f"warnings={len(self.warnings)})"
        )


def render_deck(
    plan: DeckPlan,
    pack: FormatPack,
    blueprints: BlueprintLibrary,
    out_path: Path,
    *,
    warnings: Optional[List[str]] = None,
) -> RenderResult:
    """Render ``plan`` with ``pack`` styling onto ``blueprints``; save .pptx.

    Warnings are appended to ``warnings`` (a fresh list is used when None) and
    copied onto the returned :class:`RenderResult`.
    """
    sink: List[str] = warnings if warnings is not None else []

    aspect = plan.aspect_ratio or "16:9"
    with _warnings_mod.catch_warnings(record=True) as caught:
        _warnings_mod.simplefilter("always")
        width_emu, height_emu = canvas_size(aspect)
    for record in caught:
        sink.append(str(record.message))

    canvas = Canvas(width_emu, height_emu)

    prs = Presentation()
    prs.slide_width = Emu(width_emu)
    prs.slide_height = Emu(height_emu)

    try:
        blank_layout = prs.slide_layouts[6]
    except IndexError:
        blank_layout = prs.slide_layouts[-1]

    for slide in plan.slides:
        _render_slide(prs, slide, pack, blueprints, canvas, blank_layout, aspect, sink)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    _rewrite_zip_deterministic(out_path)

    return RenderResult(
        out_path=out_path,
        slide_count=len(plan.slides),
        warnings=sink,
        plan_snapshot=plan.model_dump_json(),
    )


# --------------------------------------------------------------------------- #
# Reproducible builds
# --------------------------------------------------------------------------- #
_FIXED_ZIP_MTIME = (1980, 1, 1, 0, 0, 0)


def _rewrite_zip_deterministic(path: Path) -> None:
    """Rewrite the archive with fixed entry timestamps.

    python-pptx stamps each zip entry with the wall-clock time, so two renders
    that straddle a second boundary are not byte-identical. Normalising the
    container's timestamps makes DeckForge output reproducible: identical
    plans, packs and blueprints produce byte-identical .pptx files.
    """
    temporary = path.with_name(path.name + ".tmp")
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary, "w"
        ) as target:
            for info in source.infolist():
                fixed = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_MTIME)
                fixed.compress_type = info.compress_type
                fixed.comment = info.comment
                fixed.extra = info.extra
                target.writestr(fixed, source.read(info.filename))
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink(missing_ok=True)


# --------------------------------------------------------------------------- #
# Slide rendering
# --------------------------------------------------------------------------- #
def _render_slide(
    prs: Presentation,
    slide: SlidePlan,
    pack: FormatPack,
    blueprints: BlueprintLibrary,
    canvas: Canvas,
    blank_layout,
    aspect: str,
    sink: List[str],
) -> None:
    try:
        blueprint, effective = blueprints.resolve(slide.archetype, aspect)
    except KeyError as exc:
        sink.append(str(exc))
        _render_fallback_slide(prs, slide, blank_layout, canvas)
        return

    target = prs.slides.add_slide(blank_layout)

    known = {s.name for s in effective}
    for key in slide.slots:
        if key not in known:
            sink.append(
                f"slide {slide.n}: unused slot {key!r} (blueprint {blueprint.id!r} has no such slot)"
            )

    placements = layout_slots(slide, effective, canvas)

    # A plan-level takeaway title fills a slot named "title" when the plan
    # didn't already populate it.
    for placed in placements:
        if (
            placed.name == "title"
            and slide.title
            and (slide.slots.get("title") is None or slide.slots["title"].is_empty)
        ):
            slide.slots["title"] = SlotValue(text=slide.title)

    for placed in placements:
        value = slide.slots.get(placed.name)
        if value is None or value.is_empty:
            if not placed.slot_def.optional:
                _draw_empty_outline(target, placed, pack)
                sink.append(f"slide {slide.n}: empty required slot {placed.name}")
            continue
        _render_slot_content(target, slide, placed, value, pack, sink)

    if slide.speaker_notes:
        target.notes_slide.notes_text_frame.text = slide.speaker_notes


def _render_fallback_slide(
    prs: Presentation, slide: SlidePlan, blank_layout, canvas: Canvas
) -> None:
    target = prs.slides.add_slide(blank_layout)
    box = target.shapes.add_textbox(
        int(canvas.width_emu * 0.1),
        int(canvas.height_emu * 0.4),
        int(canvas.width_emu * 0.8),
        int(canvas.height_emu * 0.2),
    )
    box.text_frame.text = slide.title or f"[{slide.archetype}]"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string("1C1C1A")


def _render_slot_content(
    target,
    slide: SlidePlan,
    placed: PlacedSlot,
    value: SlotValue,
    pack: FormatPack,
    sink: List[str],
) -> None:
    kinds = set(placed.slot_def.kinds)

    if ContentKind.IMAGE in kinds and (value.image is not None or value.url):
        _render_image(target, placed, value, pack)
        return
    if ContentKind.CHART in kinds and value.chart is not None:
        add_chart(target, value.chart, *placed.rect, pack=pack)
        return
    if ContentKind.TABLE in kinds and value.table is not None:
        add_table(target, value.table, *placed.rect, pack=pack)
        return
    if ContentKind.DIAGRAM in kinds and value.diagram is not None:
        sink.append(
            f"slide {slide.n}: slot {placed.name}: diagram rendering deferred (placeholder)"
        )
        _draw_diagram_placeholder(target, placed, pack)
        return
    if ContentKind.PERSON in kinds and value.people:
        _render_people(target, placed, value, pack, slide.n, sink)
        return
    _render_text_slot(target, placed, value, pack, slide.n, sink)


# --------------------------------------------------------------------------- #
# Text
# --------------------------------------------------------------------------- #
def _render_text_slot(
    target,
    placed: PlacedSlot,
    value: SlotValue,
    pack: FormatPack,
    slide_n: int,
    sink: List[str],
) -> None:
    entry = _type_entry(pack, placed.slot_def.style, sink)
    bullet = ContentKind.LIST in placed.slot_def.kinds

    def _cap(text: str) -> str:
        if (entry.caps or "none").lower() == "upper":
            return text.upper()
        return text

    def _bullet(text: str) -> str:
        text = text.strip()
        if bullet and text and not text.startswith(_BULLET_CHARS):
            return f"• {text}"
        return text

    runs: List[Tuple[str, bool]] = []
    if value.number is not None:
        runs.append((_cap(str(value.number)), True))
    if value.text:
        runs.append((_cap(value.text), False))
    for paragraph in value.paragraphs or []:
        runs.append((_cap(paragraph), False))
    for item in value.items or []:
        runs.append((_cap(_bullet(item)), False))
    if value.source:
        runs.append((_cap(value.source), False))
    if not runs:
        return

    _emit_textbox(
        target,
        placed.rect,
        runs,
        entry,
        pack,
        placed.slot_def.align,
        placed.slot_def.valign,
        slide_n,
        placed.name,
        sink,
    )


def _emit_textbox(
    target,
    rect: Tuple[int, int, int, int],
    runs: Sequence[Tuple[str, bool]],
    entry: TypeScaleEntry,
    pack: FormatPack,
    align: str,
    valign: str,
    slide_n: int,
    slot_name: str,
    sink: List[str],
) -> None:
    left, top, width, height = rect
    margins = pack.style.margins
    inset_left = int(margins.left * width)
    inset_right = int(margins.right * width)
    inset_top = int(margins.top * height)
    inset_bottom = int(margins.bottom * height)
    box_width = width - inset_left - inset_right
    box_height = height - inset_top - inset_bottom

    texts = [t for t, _ in runs]
    size, overflow, _lines = fit_paragraphs(texts, entry, box_width, box_height)
    if size < entry.size_pt - 0.01:
        sink.append(
            f"slide {slide_n}: slot {slot_name}: text shrunk from {entry.size_pt}pt to {size}pt"
        )
    if overflow:
        sink.append(
            f"slide {slide_n}: slot {slot_name}: text overflows its box even at {size}pt"
        )

    textbox = target.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Emu(inset_left)
    tf.margin_right = Emu(inset_right)
    tf.margin_top = Emu(inset_top)
    tf.margin_bottom = Emu(inset_bottom)
    tf.vertical_anchor = _anchor(valign)

    paragraph_alignment = _alignment(align)
    for index, (text, is_number) in enumerate(runs):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = paragraph_alignment
        paragraph.line_spacing = entry.line_spacing
        if entry.space_after_pt:
            paragraph.space_after = Pt(entry.space_after_pt)
        run = paragraph.add_run()
        run.text = text
        run.font.name = _font_for(pack, entry)
        run.font.size = Pt(size)
        run.font.bold = entry.weight == "bold"
        run.font.color.rgb = _rgb(_text_color(pack, entry, is_number))


# --------------------------------------------------------------------------- #
# Images
# --------------------------------------------------------------------------- #
def _render_image(target, placed: PlacedSlot, value: SlotValue, pack: FormatPack) -> None:
    image = value.image
    image_path = None
    if image is not None and image.url:
        image_path = image.url
    elif value.url:
        image_path = value.url
    alt = (image.alt if image is not None else None) or "image"

    if image_path and Path(image_path).is_file():
        _place_native_image(target, placed, str(image_path), pack)
    else:
        _draw_image_placeholder(target, placed, alt, pack)


def _place_native_image(target, placed: PlacedSlot, image_path: str, pack: FormatPack) -> None:
    treatment = pack.style.image_treatment
    mode = (treatment.mode or "framed").lower()
    fit = placed.slot_def.image_fit or "crop"

    if mode == "rounded":
        shape = target.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, placed.left, placed.top, placed.width, placed.height
        )
        radius = treatment.corner_radius or pack.style.corner_radii or 0.02
        shape.adjustments[0] = max(0.0, min(0.5, radius))
        _set_picture_fill(target, shape, image_path)
        _apply_frame(shape, treatment, pack)
        return

    picture = _place_and_crop_picture(target, placed.rect, image_path, fit)
    if mode == "framed":
        _apply_frame(picture, treatment, pack)


def _place_and_crop_picture(target, rect: Tuple[int, int, int, int], image_path: str, fit: str):
    left, top, width, height = rect
    picture = target.shapes.add_picture(image_path, left, top)
    natural_w, natural_h = picture.width, picture.height
    if fit == "crop":
        scale = max(width / natural_w, height / natural_h)
    else:  # "fit": letterbox inside the box
        scale = min(width / natural_w, height / natural_h)
    scaled_w = int(round(natural_w * scale))
    scaled_h = int(round(natural_h * scale))
    picture.left = left + (width - scaled_w) // 2
    picture.top = top + (height - scaled_h) // 2
    picture.width = scaled_w
    picture.height = scaled_h
    if fit == "crop":
        if scaled_w > width:
            frac = min(1.0, (scaled_w - width) / scaled_w)
            picture.crop_left = frac / 2
            picture.crop_right = frac / 2
        if scaled_h > height:
            frac = min(1.0, (scaled_h - height) / scaled_h)
            picture.crop_top = frac / 2
            picture.crop_bottom = frac / 2
    return picture


def _apply_frame(shape, treatment, pack: FormatPack) -> None:
    frame_color = treatment.frame_color or _hex(pack, "neutral1", "muted-text", "line", "text")
    weight = treatment.frame_weight_pt or pack.style.line_weight_pt or 1.0
    shape.line.color.rgb = _rgb(frame_color)
    shape.line.width = Pt(weight)


def _set_picture_fill(target, shape, image_path: str) -> None:
    image_part, r_id = target.part.get_or_add_image_part(image_path)
    sp_pr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:grpFill"):
        for element in sp_pr.findall(qn(tag)):
            sp_pr.remove(element)
    blip_fill = etree.SubElement(sp_pr, qn("a:blipFill"))
    blip = etree.SubElement(blip_fill, qn("a:blip"))
    blip.set(qn("r:embed"), r_id)
    stretch = etree.SubElement(blip_fill, qn("a:stretch"))
    etree.SubElement(stretch, qn("a:fillRect"))
    geometry = sp_pr.find(qn("a:prstGeom"))
    geometry.addnext(blip_fill)


def _draw_image_placeholder(target, placed: PlacedSlot, alt: str, pack: FormatPack) -> None:
    outline = target.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, placed.left, placed.top, placed.width, placed.height
    )
    outline.fill.background()
    outline.line.color.rgb = _rgb(_hex(pack, "neutral1", "muted-text", "line", "text"))
    outline.line.width = Pt(1)

    label = target.shapes.add_textbox(
        placed.left + 1, placed.top + 1, max(1, placed.width - 2), max(1, placed.height - 2)
    )
    label.text_frame.word_wrap = True
    paragraph = label.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = f"[image placeholder: {alt}]"
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(_hex(pack, "muted-text", "text"))


# --------------------------------------------------------------------------- #
# People cards
# --------------------------------------------------------------------------- #
def _render_people(
    target,
    placed: PlacedSlot,
    value: SlotValue,
    pack: FormatPack,
    slide_n: int,
    sink: List[str],
) -> None:
    people = [p for p in (value.people or []) if p.name]
    if not people:
        return
    entry = _type_entry(pack, placed.slot_def.style, sink)
    role_entry = _type_entry(pack, _role_style_name(entry.name), sink)

    gap = int(placed.width * 0.02)
    card_w = (placed.width - gap * (len(people) - 1)) // len(people)
    for index, person in enumerate(people):
        left = placed.left + index * (card_w + gap)
        portrait_h = 0
        if person.image_url and Path(person.image_url).is_file():
            portrait_h = min(card_w, int(placed.height * 0.35))
            picture = _place_and_crop_picture(
                target,
                (
                    left + (card_w - portrait_h) // 2,
                    placed.top,
                    portrait_h,
                    portrait_h,
                ),
                person.image_url,
                "crop",
            )
            picture.line.fill.background()
        name_top = placed.top + portrait_h
        name_h = max(1, int((placed.height - portrait_h) * 0.6))
        _emit_textbox(
            target,
            (left, name_top, card_w, name_h),
            [(person.name, False)],
            entry,
            pack,
            "center",
            "middle",
            slide_n,
            placed.name,
            sink,
        )
        role_top = name_top + name_h
        role_h = max(1, placed.height - portrait_h - name_h)
        _emit_textbox(
            target,
            (left, role_top, card_w, role_h),
            [(person.role, False)],
            role_entry,
            pack,
            "center",
            "top",
            slide_n,
            placed.name,
            sink,
        )


def _role_style_name(entry_name: str) -> str:
    return "caption" if entry_name in _HEADING_STYLES else "small"


# --------------------------------------------------------------------------- #
# Empty / placeholder outlines
# --------------------------------------------------------------------------- #
def _draw_empty_outline(target, placed: PlacedSlot, pack: FormatPack) -> None:
    outline = target.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, placed.left, placed.top, placed.width, placed.height
    )
    outline.fill.background()
    outline.line.color.rgb = _rgb(_hex(pack, "neutral1", "muted-text", "line", "text"))
    outline.line.width = Pt(0.75)
    label = target.shapes.add_textbox(
        placed.left + 1, placed.top + 1, max(1, placed.width - 2), max(1, placed.height - 2)
    )
    paragraph = label.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = f"[empty: {placed.name}]"
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(_hex(pack, "muted-text", "text"))


def _draw_diagram_placeholder(target, placed: PlacedSlot, pack: FormatPack) -> None:
    _draw_empty_outline(target, placed, pack)


# --------------------------------------------------------------------------- #
# Style lookups (never raise; fall back + warn)
# --------------------------------------------------------------------------- #
def _type_entry(pack: FormatPack, name: str, sink: List[str]) -> TypeScaleEntry:
    scale = pack.style.type_scale
    try:
        return scale.entry(name)
    except KeyError:
        pass
    sink.append(f"style {name!r} not in type scale; falling back to 'body'")
    try:
        return scale.entry("body")
    except KeyError:
        return TypeScaleEntry(name="body", size_pt=16, line_spacing=1.2)


def _hex(pack: FormatPack, *roles: str) -> str:
    palette = pack.style.palette
    for role in roles:
        try:
            return palette.hex(role)
        except KeyError:
            continue
    return palette.colors[0].hex


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#"))


def _text_color(pack: FormatPack, entry: TypeScaleEntry, is_number: bool) -> str:
    if is_number or entry.name in ("number", "big-number", "kicker"):
        return _hex(pack, "accent1")
    if entry.name in ("caption", "small"):
        return _hex(pack, "muted-text")
    return _hex(pack, "text")


def _font_for(pack: FormatPack, entry: TypeScaleEntry) -> str:
    fonts = pack.style.fonts
    return fonts.heading if entry.name in _HEADING_STYLES else fonts.body


def _anchor(valign: str):
    return {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }.get((valign or "top").lower(), MSO_ANCHOR.TOP)


def _alignment(align: str):
    return {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get((align or "left").lower(), PP_ALIGN.LEFT)
