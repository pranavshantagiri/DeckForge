"""Native python-pptx tables styled strictly from the pack's StyleProfile.

Clean look: header row filled with the accent, zebra-striped body with a
neutral tint, body font at small size. No gridlines are drawn — fill carries
the structure, matching the anti-"AI-look" rule.
"""

from __future__ import annotations

from typing import List

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from deckforge_core.schemas.deck_plan import TableSpec
from deckforge_core.schemas.pack import FormatPack

_CELL_MARGIN_EMU = 45720  # 0.05"


def _hex(pack: FormatPack, *roles: str) -> str:
    palette = pack.style.palette
    for role in roles:
        try:
            return palette.hex(role)
        except KeyError:
            continue
    return palette.colors[0].hex


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#"))


def _blend(hex_a: str, hex_b: str, t: float) -> str:
    """Linearly blend two hex colours; ``t=0`` -> ``a``, ``t=1`` -> ``b``."""
    a = [int(hex_a[i : i + 2], 16) for i in (1, 3, 5)]
    b = [int(hex_b[i : i + 2], 16) for i in (1, 3, 5)]
    out = [round(aa + (bb - aa) * t) for aa, bb in zip(a, b)]
    return "#{:02X}{:02X}{:02X}".format(*out)


def _col_widths(spec: TableSpec, total: int, columns: int) -> List[int]:
    if spec.col_widths:
        fractions = list(spec.col_widths)
    else:
        fractions = [1.0 / columns] * columns
    if len(fractions) != columns or sum(fractions) <= 0:
        fractions = [1.0 / columns] * columns
    widths = [int(total * f / sum(fractions)) for f in fractions]
    widths[-1] = total - sum(widths[:-1])  # absorb rounding
    return widths


def add_table(
    slide,
    spec: TableSpec,
    left: int,
    top: int,
    width: int,
    height: int,
    pack: FormatPack,
):
    """Draw ``spec`` as a native table on ``slide``; return the table object."""
    columns = len(spec.columns)
    header = bool(spec.header_row)
    n_rows = len(spec.rows) + (1 if header else 0)
    n_rows = max(n_rows, 1)

    graphic_frame = slide.shapes.add_table(n_rows, columns, left, top, width, height)
    table = graphic_frame.table
    table.first_row = False
    table.horz_banding = False

    col_widths = _col_widths(spec, width, columns)
    for i, cw in enumerate(col_widths):
        table.columns[i].width = cw
    row_height = max(1, height // n_rows)
    for i in range(n_rows):
        table.rows[i].height = row_height

    style = pack.style
    body_font = style.fonts.body
    accent_hex = _hex(pack, "accent1")
    body_hex = _hex(pack, "text")
    zebra_hex = _hex(pack, "card-bg", "neutral1", "neutral2", "bg")
    zebra_alt = _blend(_hex(pack, "accent1", "text"), zebra_hex, 0.92)

    small_pt = _small_pt(pack)

    def _style_cell(row: int, col: int, text: str, is_header: bool, is_zebra: bool) -> None:
        cell = table.cell(row, col)
        cell.margin_left = Emu(_CELL_MARGIN_EMU)
        cell.margin_right = Emu(_CELL_MARGIN_EMU)
        cell.margin_top = Emu(27432)
        cell.margin_bottom = Emu(27432)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if is_header:
            cell.fill.fore_color.rgb = _rgb(accent_hex)
        elif is_zebra:
            cell.fill.fore_color.rgb = _rgb(zebra_alt)
        else:
            cell.fill.fore_color.rgb = _rgb(zebra_hex)

        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = text
        run.font.name = body_font
        run.font.size = Pt(small_pt)
        run.font.bold = is_header
        run.font.color.rgb = _rgb("#FFFFFF") if is_header else _rgb(body_hex)

    for col in range(columns):
        header_text = spec.columns[col] if col < len(spec.columns) else ""
        _style_cell(0, col, header_text, is_header=header, is_zebra=False)

    for r, row in enumerate(spec.rows):
        target = r + (1 if header else 0)
        zebra = bool(spec.zebra_rows) and (r % 2 == 1)
        for col in range(columns):
            text = row[col] if col < len(row) else ""
            _style_cell(target, col, text, is_header=False, is_zebra=zebra)

    return table


def _small_pt(pack: FormatPack) -> float:
    for name in ("small", "caption", "body"):
        try:
            return pack.style.type_scale.entry(name).size_pt
        except KeyError:
            continue
    return 12.0
